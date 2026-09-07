# -*- coding: utf-8 -*-
"""学习资料真实文件生成模块

根据资料 ID 动态生成真实可下载的文件（PDF / DOCX），内容为真实学科笔记/模板。
- id=1 高等数学(下)期末复习笔记  -> PDF
- id=2 GIS空间分析实验报告模板   -> DOCX
- id=3 大学英语四级真题及答案    -> PDF
- id=4 数据结构期末重点整理      -> PDF
"""
from io import BytesIO

# ---------- PDF（reportlab，内置中文字体） ----------
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import (HRFlowable, PageBreak, Paragraph, SimpleDocTemplate,
                                Spacer)

pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
_ZH = "STSong-Light"

_styles = getSampleStyleSheet()
_TITLE = ParagraphStyle("zh-title", parent=_styles["Title"], fontName=_ZH,
                        fontSize=20, leading=28, spaceAfter=12)
_H1 = ParagraphStyle("zh-h1", parent=_styles["Heading1"], fontName=_ZH,
                     fontSize=16, leading=22, textColor=colors.HexColor("#1d6df0"),
                     spaceBefore=14, spaceAfter=8)
_H2 = ParagraphStyle("zh-h2", parent=_styles["Heading2"], fontName=_ZH,
                     fontSize=13, leading=18, spaceBefore=8, spaceAfter=4)
_BODY = ParagraphStyle("zh-body", parent=_styles["BodyText"], fontName=_ZH,
                       fontSize=10.5, leading=16, spaceAfter=4)
_NOTE = ParagraphStyle("zh-note", parent=_BODY, textColor=colors.HexColor("#666"),
                       fontSize=9.5, leftIndent=10)


def _make_pdf(title: str, sections: list) -> bytes:
    """sections: [(h1, [(h2, [para...])...]), ...]"""
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=22 * mm, rightMargin=22 * mm,
                            topMargin=20 * mm, bottomMargin=18 * mm,
                            title=title)
    story = [Paragraph(title, _TITLE),
             HRFlowable(width="100%", thickness=1, color=colors.HexColor("#1d6df0")),
             Spacer(1, 8)]
    for h1, subs in sections:
        story.append(Paragraph(h1, _H1))
        for h2, paras in subs:
            story.append(Paragraph(h2, _H2))
            for p in paras:
                story.append(Paragraph(p, _BODY))
    doc.build(story)
    return buf.getvalue()


# ---------- DOCX（python-docx） ----------
def _make_docx(title: str, lines: list) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = None  # 默认字号

    h = doc.add_heading(title, level=0)
    for run in h.runs:
        run.font.name = "宋体"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for kind, text in lines:
        if kind == "h2":
            p = doc.add_heading(text, level=2)
        elif kind == "line":
            doc.add_paragraph(text)
        elif kind == "empty":
            doc.add_paragraph("")
        elif kind == "fill":
            doc.add_paragraph("＿" * 36)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- 各资料内容 ----------

def _gaoshu_pdf() -> bytes:
    sections = [
        ("一、函数、极限与连续", [
            ("1.1 核心概念", [
                "• 极限定义：当 x→x₀ 时，f(x) 无限趋近于 A，记为 lim f(x) = A。",
                "• 两个重要极限：lim(sinx/x)=1 (x→0)；lim(1+1/n)ⁿ = e (n→∞)。",
                "• 连续条件：f(x₀) 有定义、极限存在、极限值等于函数值。",
            ]),
            ("1.2 常考题型", [
                "【例】求 lim(x→0)(1-cosx)/x²。",
                "【解】1-cosx ~ x²/2，故原式 = lim (x²/2)/x² = 1/2。",
            ]),
        ]),
        ("二、导数与微分", [
            ("2.1 基本求导公式", [
                "• (xⁿ)' = nxⁿ⁻¹； (sinx)' = cosx； (cosx)' = -sinx；",
                "• (eˣ)' = eˣ； (lnx)' = 1/x； 复合函数链式法则 (f(g(x)))' = f'(g(x))·g'(x)。",
            ]),
            ("2.2 常考题型", [
                "【例】y = e^(2x)·sinx，求 y'。",
                "【解】y' = 2e^(2x)·sinx + e^(2x)·cosx = e^(2x)(2sinx + cosx)。",
            ]),
        ]),
        ("三、不定积分与定积分", [
            ("3.1 基本积分公式", [
                "• ∫xⁿdx = xⁿ⁺¹/(n+1) + C (n≠-1)； ∫(1/x)dx = ln|x| + C；",
                "• ∫sinx dx = -cosx + C； ∫cosx dx = sinx + C；",
                "• 分部积分：∫u dv = uv - ∫v du； 换元积分：凑微分与三角代换。",
            ]),
            ("3.2 定积分应用", [
                "• 曲边梯形面积：S = ∫ₐᵇ|f(x)|dx；",
                "• 旋转体体积：V = π∫ₐᵇf²(x)dx。",
            ]),
        ]),
        ("四、多元函数微分学", [
            ("4.1 偏导数与全微分", [
                "• z=f(x,y) 对 x 偏导：把 y 当常数对 x 求导；",
                "• 全微分：dz = fₓdx + fᵧdy。",
            ]),
            ("4.2 常考题型", [
                "【例】z = x²y + sin(xy)，求 ∂z/∂x。",
                "【解】∂z/∂x = 2xy + y·cos(xy)。",
            ]),
        ]),
        ("五、无穷级数", [
            ("5.1 判敛方法", [
                "• 正项级数：比较审敛、比值审敛（lim aₙ₊₁/aₙ < 1 收敛）；",
                "• 交错级数：莱布尼茨判别法（单调递减趋于 0 则收敛）。",
            ]),
            ("5.2 幂级数", [
                "• 收敛半径 R = lim|aₙ/aₙ₊₁|；",
                "• 常见展开：eˣ = Σxⁿ/n!； ln(1+x) = Σ(-1)ⁿ⁺¹xⁿ/n。",
            ]),
        ]),
    ]
    return _make_pdf("高等数学（下）期末复习笔记", sections)


def _gis_docx() -> bytes:
    lines = [
        ("h2", "一、基本信息"),
        ("line", "实验名称：GIS 空间分析实验"),
        ("line", "课程名称：地理信息系统原理与应用"),
        ("line", "实验人：＿＿＿＿＿＿  学号：＿＿＿＿＿＿  班级：＿＿＿＿＿＿"),
        ("line", "实验日期：＿＿＿＿年＿＿月＿＿日  指导教师：＿＿＿＿＿＿"),
        ("empty", ""),
        ("h2", "二、实验目的"),
        ("line", "1. 掌握 GIS 空间分析的基本原理与操作流程；"),
        ("line", "2. 熟练使用缓冲区分析、叠加分析、网络分析等空间分析方法；"),
        ("line", "3. 能够针对实际问题设计并完成空间分析方案，正确解读分析结果。"),
        ("empty", ""),
        ("h2", "三、实验环境与数据"),
        ("line", "实验平台：ArcGIS / QGIS / 自研 WebGIS 平台"),
        ("line", "实验数据：＿＿＿＿＿＿（数据来源、坐标系、比例尺等）"),
        ("empty", ""),
        ("h2", "四、实验原理"),
        ("line", "（简述本次实验所涉及的空间分析方法的原理，例如缓冲区分析以指定要素为中心、"),
        ("line", "按给定半径生成一定范围的多边形，用于分析要素的影响范围。）"),
        ("empty", ""),
        ("h2", "五、实验步骤"),
        ("line", "1. 数据加载与坐标系统一；"),
        ("line", "2. ＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿；"),
        ("line", "3. ＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿；"),
        ("line", "4. 结果输出与制图表达。"),
        ("empty", ""),
        ("h2", "六、实验结果与分析"),
        ("line", "（插入实验截图，并对结果进行定量与定性分析，说明其反映的空间规律。）"),
        ("empty", ""),
        ("h2", "七、实验结论与心得"),
        ("line", "（总结实验收获，指出实验中遇到的问题及解决方法。）"),
        ("empty", ""),
        ("h2", "八、评分"),
        ("line", "成绩：＿＿＿＿  教师评语：＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿"),
    ]
    return _make_docx("GIS 空间分析实验报告模板", lines)


def _cet4_pdf() -> bytes:
    sections = [
        ("一、写作（Writing, 106.5 分）", [
            ("题目示例", [
                "Directions: For this part, you are allowed 30 minutes to write a short",
                "essay on the topic: The Importance of Environmental Protection.",
                "You should write at least 120 words but no more than 180 words.",
            ]),
            ("参考范文（节选）", [
                "Environmental protection has become an increasingly significant issue in",
                "modern society. With the rapid development of industry, our natural",
                "environment is facing unprecedented challenges...",
            ]),
        ]),
        ("二、听力理解（Listening, 248.5 分）", [
            ("Section A 短篇新闻", [
                "1. A) It was launched by a charity.  B) It attracted many volunteers.",
                "C) It helped the local farmers.    D) It lasted for a whole week.",
            ]),
            ("答题技巧", [
                "• 听前快速浏览选项，预测问题方向；",
                "• 注意数字、地点、人名等关键词；",
                "• 同义替换是常见考点。",
            ]),
        ]),
        ("三、阅读理解（Reading, 248.5 分）", [
            ("选词填空（节选）", [
                "The Internet has changed the way we ___ (live) and work.",
                "【答案】live",
            ]),
            ("仔细阅读", [
                "Q: What is the main purpose of the passage?",
                "A) To introduce a new technology.",
                "B) To discuss the effects of climate change.",
                "C) To call for public participation.",
                "D) To compare different countries.",
                "【解析】主旨题：通读首尾段，文章呼吁公众参与环境保护，故选 C。",
            ]),
        ]),
        ("四、翻译（Translation, 106.5 分）", [
            ("汉译英示例", [
                "原文：中国是一个拥有悠久历史和灿烂文化的国家。",
                "译文：China is a country with a long history and splendid culture.",
                "得分要点：主谓宾完整、时态一致、用词准确。",
            ]),
        ]),
        ("五、备考建议", [
            ("词汇与真题", [
                "• 每日背诵高频词汇 50 个，滚动复习；",
                "• 近 3 年真题精做精析，错题整理成册；",
                "• 听力坚持精听训练，阅读限时练习。",
            ]),
        ]),
    ]
    return _make_pdf("大学英语四级真题精选与解析", sections)


def _datastruct_pdf() -> bytes:
    sections = [
        ("一、线性表", [
            ("1.1 重点", [
                "• 顺序表：随机存取，插入删除需移动元素（平均移动 n/2 个）；",
                "• 链表：顺序存取，插入删除 O(1)（已知位置）；",
                "• 头结点作用：统一空表与非空表的操作。",
            ]),
        ]),
        ("二、栈与队列", [
            ("2.1 重点", [
                "• 栈：后进先出 LIFO，应用——括号匹配、表达式求值、函数调用；",
                "• 队列：先进先出 FIFO，循环队列判满 (rear+1)%MaxSize == front；",
                "• 共享栈、双端队列了解即可。",
            ]),
        ]),
        ("三、树与二叉树", [
            ("3.1 二叉树性质", [
                "• 第 i 层最多 2^(i-1) 个结点；深度为 k 最多 2^k - 1 个结点；",
                "• n₀ = n₂ + 1（叶子数 = 度为 2 的结点数 + 1）；",
                "• 完全二叉树用数组存储，结点 i 的孩子为 2i、2i+1。",
            ]),
            ("3.2 遍历", [
                "• 前序/中序/后序/层序四种遍历，已知前序+中序可唯一确定二叉树。",
            ]),
        ]),
        ("四、图", [
            ("4.1 重点", [
                "• 存储：邻接矩阵（稠密）、邻接表（稀疏）；",
                "• 遍历：DFS（栈/递归）、BFS（队列）；",
                "• 最小生成树：Prim（加点）、Kruskal（加边）；",
                "• 最短路径：Dijkstra（单源，非负权）、Floyd（多源）；",
                "• 拓扑排序：AOV 网，环检测。",
            ]),
        ]),
        ("五、查找", [
            ("5.1 重点", [
                "• 顺序查找 ASL=(n+1)/2；折半查找 ASL=log₂(n+1)-1（有序）；",
                "• 二叉排序树 BST：中序有序；平衡树 AVL 旋转调整；",
                "• 哈希表：除留余数法、线性探测、链地址法，装填因子 α。",
            ]),
        ]),
        ("六、排序", [
            ("6.1 常考排序", [
                "• 快排 O(nlogn) 平均 / O(n²) 最坏，不稳定；",
                "• 归并 O(nlogn)，稳定，空间 O(n)；",
                "• 堆排序 O(nlogn)，不稳定，建堆 O(n)；",
                "• 简单排序对比表：冒泡/插入/选择 O(n²)；希尔 O(n^1.3)。",
            ]),
        ]),
    ]
    return _make_pdf("数据结构期末重点整理", sections)


_GENERATORS = {
    1: _gaoshu_pdf,
    2: _gis_docx,
    3: _cet4_pdf,
    4: _datastruct_pdf,
}


def generate_material_file(mid: int, title: str):
    """返回 (bytes, filename, media_type)；未知类型抛 ValueError"""
    if mid not in _GENERATORS:
        raise ValueError(f"资料 {mid} 暂无下载文件")
    data = _GENERATORS[mid]()
    if mid == 2:
        filename = f"{title}.docx"
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        filename = f"{title}.pdf"
        media_type = "application/pdf"
    return data, filename, media_type
